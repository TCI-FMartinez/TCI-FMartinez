import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from modulos.param_core import (
    ParamDocument,
    ParamEditorError,
    build_export_plan,
    create_workbook_for_factory,
    find_factory_dirs,
    load_metadata,
    validate_collisions,
    write_export_plan,
)
from modulos.param_methadata import metadata_path


SAMPLE_A = (
    b"# comentario que debe conservarse\r\n"
    b"N001RACERO AL CARBONO\r\n"
    b"N004R5.00\r\n"
    b"N005R6000\r\n"
    b"N026RSTD\r\n"
    b"N026RSEGUNDA\r\n"
    b"N037RO2\r\n"
)

SAMPLE_B = (
    b"N001SINOXIDABLE\n"
    b"N004R2.0\n"
    b"N005R6000\n"
    b"N037RN2\n"
)


class ParamEditorTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.process_root = self.root / "para_procesar"
        self.factory = self.process_root / "MAQUINA" / "PARAMS_LASER" / "factory"
        self.material = self.factory / "STEEL"
        self.material.mkdir(parents=True)
        (self.material / "receta_a").write_bytes(SAMPLE_A)
        (self.material / "receta_b").write_bytes(SAMPLE_B)
        self.metadata = load_metadata(metadata_path())

    def tearDown(self):
        self.tmp.cleanup()

    def test_document_no_edits_is_byte_exact(self):
        doc = ParamDocument.from_bytes(SAMPLE_A)
        self.assertEqual(doc.render({}, {}), SAMPLE_A)

    def test_duplicate_occurrences_are_addressable(self):
        doc = ParamDocument.from_bytes(SAMPLE_A)
        values = doc.value_map()
        self.assertEqual(values["N026"], "STD")
        self.assertEqual(values["N026#2"], "SEGUNDA")
        out = doc.render({"N026#2": "EDITADA"}, {"N026": "R"})
        self.assertIn(b"N026RSTD\r\nN026REDITADA\r\n", out)
        self.assertIn(b"# comentario que debe conservarse\r\n", out)

    def test_recursive_factory_discovery(self):
        direct = self.process_root / "OTRA" / "factory"
        direct.mkdir(parents=True)
        found = find_factory_dirs(self.process_root)
        self.assertIn(self.factory, found)
        self.assertIn(direct, found)

    def test_roundtrip_workbook_without_edits_preserves_bytes(self):
        xlsx, count = create_workbook_for_factory(self.factory, self.process_root, self.metadata)
        self.assertEqual(count, 2)
        output = self.root / "out"
        plan = build_export_plan(xlsx, self.process_root, output, self.metadata)
        validate_collisions(plan)
        write_export_plan(plan, output)
        self.assertEqual((output / "MAQUINA/PARAMS_LASER/factory/STEEL/receta_a.lparam").read_bytes(), SAMPLE_A)
        self.assertEqual((output / "MAQUINA/PARAMS_LASER/factory/STEEL/receta_b.lparam").read_bytes(), SAMPLE_B)

    def test_single_edit_preserves_status_and_other_lines(self):
        xlsx, _ = create_workbook_for_factory(self.factory, self.process_root, self.metadata)
        wb = load_workbook(xlsx)
        ws = wb["Hoja1"]
        headers = {cell.value: cell.column for cell in ws[1]}
        ws.cell(2, headers["N004"], "6.50")
        wb.save(xlsx)

        output = self.root / "out"
        plan = build_export_plan(xlsx, self.process_root, output, self.metadata)
        write_export_plan(plan, output)
        out = (output / "MAQUINA/PARAMS_LASER/factory/STEEL/receta_a.lparam").read_bytes()
        self.assertIn(b"N004R6.50\r\n", out)
        self.assertIn(b"N001RACERO AL CARBONO\r\n", out)
        self.assertIn(b"N026RSEGUNDA\r\n", out)

    def test_invalid_edited_thickness_aborts(self):
        xlsx, _ = create_workbook_for_factory(self.factory, self.process_root, self.metadata)
        wb = load_workbook(xlsx)
        ws = wb["Hoja1"]
        headers = {cell.value: cell.column for cell in ws[1]}
        ws.cell(2, headers["N004"], "-1")
        wb.save(xlsx)
        with self.assertRaises(ParamEditorError):
            build_export_plan(xlsx, self.process_root, self.root / "out", self.metadata)

    def test_output_collision_aborts_before_write(self):
        xlsx, _ = create_workbook_for_factory(self.factory, self.process_root, self.metadata)
        wb = load_workbook(xlsx)
        ws = wb["Hoja1"]
        headers = {cell.value: cell.column for cell in ws[1]}
        ws.cell(2, headers["__archivo_salida"], "MISMO.lparam")
        ws.cell(3, headers["__archivo_salida"], "mismo.LPARAM")
        wb.save(xlsx)
        plan = build_export_plan(xlsx, self.process_root, self.root / "out", self.metadata)
        with self.assertRaises(ParamEditorError):
            validate_collisions(plan)
        self.assertFalse((self.root / "out").exists())

    def test_control_path_cannot_be_changed(self):
        xlsx, _ = create_workbook_for_factory(self.factory, self.process_root, self.metadata)
        wb = load_workbook(xlsx)
        ws = wb["Hoja1"]
        headers = {cell.value: cell.column for cell in ws[1]}
        ws.cell(2, headers["N000"], "../OTRA")
        wb.save(xlsx)
        with self.assertRaises(ParamEditorError):
            build_export_plan(xlsx, self.process_root, self.root / "out", self.metadata)


if __name__ == "__main__":
    unittest.main()
