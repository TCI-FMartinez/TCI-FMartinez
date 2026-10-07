from __future__ import annotations

import base64
import csv
import json
import os
import re
import shutil
import tempfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

PARAM_RE = re.compile(r"^(N\d{3})([RS])(.*)$")
COLUMN_RE = re.compile(r"^(N\d{3})(?:#(\d+))?$")
EXCLUDED_SUFFIXES = {
    ".ini", ".txt", ".spec", ".exe", ".csv", ".xls", ".xlsx",
    ".zip", ".7z", ".rar", ".png", ".jpg", ".jpeg", ".log", ".rd",
}
SOURCE_SHEET = "_SOURCE"
VISIBLE_SHEET = "Hoja1"
SPECIAL_HEADERS = ["N000", "__subcarpeta", "__archivo_origen", "__archivo_salida", "__id"]


class ParamEditorError(RuntimeError):
    pass


@dataclass(frozen=True)
class ParamEntry:
    key: str
    status: str
    value: str
    occurrence: int
    line_index: int

    @property
    def column(self) -> str:
        return self.key if self.occurrence == 1 else f"{self.key}#{self.occurrence}"


@dataclass
class ParamDocument:
    raw_bytes: bytes
    encoding: str
    entries: list[ParamEntry]
    newline: str
    had_final_newline: bool

    @classmethod
    def from_bytes(cls, raw: bytes) -> "ParamDocument":
        encoding = "utf-8"
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            encoding = "cp1252"
            text = raw.decode("cp1252")

        if "\r\n" in text:
            newline = "\r\n"
        elif "\n" in text:
            newline = "\n"
        elif "\r" in text:
            newline = "\r"
        else:
            newline = os.linesep

        had_final_newline = text.endswith(("\r\n", "\n", "\r"))
        logical_lines = text.splitlines()
        counts: Counter[str] = Counter()
        entries: list[ParamEntry] = []
        for idx, line in enumerate(logical_lines):
            m = PARAM_RE.match(line)
            if not m:
                continue
            key, status, value = m.groups()
            counts[key] += 1
            entries.append(ParamEntry(key, status, value, counts[key], idx))
        return cls(raw, encoding, entries, newline, had_final_newline)

    @classmethod
    def from_path(cls, path: Path) -> "ParamDocument":
        return cls.from_bytes(path.read_bytes())

    def value_map(self) -> dict[str, str]:
        return {entry.column: entry.value for entry in self.entries}

    def status_map(self) -> dict[str, str]:
        return {entry.column: entry.status for entry in self.entries}

    def render(self, edits: dict[str, str], default_status: dict[str, str]) -> bytes:
        if not edits:
            return self.raw_bytes

        text = self.raw_bytes.decode(self.encoding)
        lines = text.splitlines()
        entry_by_column = {entry.column: entry for entry in self.entries}

        additions: list[str] = []
        for column, new_value in edits.items():
            m = COLUMN_RE.match(column)
            if not m:
                continue
            if column in entry_by_column:
                entry = entry_by_column[column]
                lines[entry.line_index] = f"{entry.key}{entry.status}{new_value}"
            else:
                key = m.group(1)
                occurrence = int(m.group(2) or "1")
                existing_count = sum(1 for e in self.entries if e.key == key)
                if new_value != "" and occurrence == existing_count + 1:
                    status = default_status.get(key, "R")
                    additions.append(f"{key}{status}{new_value}")
                elif new_value != "":
                    raise ParamEditorError(
                        f"No se puede crear {column}: falta una ocurrencia anterior de {key}."
                    )

        if additions:
            lines.extend(additions)

        rendered = self.newline.join(lines)
        if self.had_final_newline or additions:
            rendered += self.newline
        return rendered.encode(self.encoding)


def column_base_key(column: str) -> str | None:
    m = COLUMN_RE.match(column)
    return m.group(1) if m else None


def safe_relative_path(value: str, field: str) -> Path:
    p = Path(str(value).replace("\\", "/"))
    if p.is_absolute() or any(part in ("", ".", "..") for part in p.parts):
        raise ParamEditorError(f"Ruta no válida en {field}: {value!r}")
    return p


def normalize_output_name(value: str) -> str:
    name = str(value).strip()
    if not name:
        raise ParamEditorError("El nombre de archivo de salida está vacío.")
    if "/" in name or "\\" in name or name in {".", ".."}:
        raise ParamEditorError(f"Nombre de archivo de salida no válido: {name!r}")
    if not name.lower().endswith(".lparam"):
        name += ".lparam"
    return name


def is_parameter_file(path: Path) -> bool:
    return path.is_file() and path.name != "parametros.xlsx" and path.suffix.lower() not in EXCLUDED_SUFFIXES


def find_factory_dirs(root: Path) -> list[Path]:
    root = Path(root)
    if not root.exists():
        return []
    found = [p for p in root.rglob("factory") if p.is_dir()]
    return sorted(found, key=lambda p: str(p).casefold())


def find_parameter_files(factory_dir: Path) -> list[Path]:
    files = [p for p in factory_dir.rglob("*") if is_parameter_file(p)]
    return sorted(files, key=lambda p: str(p.relative_to(factory_dir)).casefold())


def load_metadata(metadata_path: Path) -> dict[str, dict[str, object]]:
    """Carga todas las hojas de methadata.xlsx tolerando columnas ausentes o renombradas."""
    metadata_path = Path(metadata_path)
    if not metadata_path.exists():
        return {}

    wb = load_workbook(metadata_path, read_only=True, data_only=True)
    result: dict[str, dict[str, object]] = {}
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            continue
        headers = {str(v).strip().lower(): i for i, v in enumerate(header_row) if v is not None}

        def index_of(*names: str) -> int | None:
            for name in names:
                idx = headers.get(name.lower())
                if idx is not None:
                    return idx
            return None

        key_idx = index_of("properties", "property", "key")
        status_idx = index_of("r=escritura/s=solo lectura", "r/s", "status")
        title_idx = index_of("titulo de columna", "título de columna", "titulo", "title")
        type_idx = index_of("tipo de dato", "type")
        max_idx = index_of("valor max.", "valor max", "max")
        min_idx = index_of("valor min.", "valor min", "min")
        if key_idx is None:
            continue

        for row in rows:
            if key_idx >= len(row) or row[key_idx] is None:
                continue
            key = str(row[key_idx]).strip()
            if not re.fullmatch(r"N\d{3}", key):
                continue

            def get(idx: int | None):
                return row[idx] if idx is not None and idx < len(row) else None

            result[key] = {
                "status": (str(get(status_idx)).strip().upper() if get(status_idx) is not None else "R"),
                "title": (str(get(title_idx)).strip() if get(title_idx) is not None else ""),
                "data_type": (str(get(type_idx)).strip().lower() if get(type_idx) is not None else ""),
                "max": get(max_idx),
                "min": get(min_idx),
                "sheet": ws.title,
            }
    return result


def metadata_default_status(metadata: dict[str, dict[str, object]]) -> dict[str, str]:
    return {
        key: (str(info.get("status") or "R").upper() if str(info.get("status") or "R").upper() in {"R", "S"} else "R")
        for key, info in metadata.items()
    }


def validate_edit(key: str, value: str, metadata: dict[str, dict[str, object]]) -> list[str]:
    """Valida solo valores editados; no invalida datos históricos no modificados."""
    errors: list[str] = []
    info = metadata.get(key, {})
    title = str(info.get("title") or "").lower()
    data_type = str(info.get("data_type") or "").lower()
    value = value.strip()

    if value == "":
        return errors

    numeric_type = data_type in {"dist", "feedrate", "power", "frequency", "duty", "press", "focal", "float", "int", "integer", "number"}
    looks_duty = "duty" in title
    must_numeric = numeric_type or key in {"N004", "N005"} or looks_duty

    number: float | None = None
    if must_numeric:
        try:
            number = float(value.replace(",", "."))
        except ValueError:
            errors.append(f"{key} debe ser numérico y vale {value!r}.")
            return errors

    if key == "N004" and number is not None and number <= 0:
        errors.append("N004 (espesor) debe ser mayor que 0.")
    if key == "N005" and number is not None and number < 0:
        errors.append("N005 (potencia) no puede ser negativa.")
    if (data_type == "duty" or looks_duty) and number is not None and not 0 <= number <= 100:
        errors.append(f"{key} (duty) debe estar entre 0 y 100.")

    min_value = info.get("min")
    max_value = info.get("max")
    if number is not None and min_value not in (None, ""):
        try:
            if number < float(min_value):
                errors.append(f"{key}={number} es menor que el mínimo {min_value}.")
        except (TypeError, ValueError):
            pass
    if number is not None and max_value not in (None, ""):
        try:
            if number > float(max_value):
                errors.append(f"{key}={number} supera el máximo {max_value}.")
        except (TypeError, ValueError):
            pass
    return errors


def _header_comment(column: str, metadata: dict[str, dict[str, object]]) -> str:
    if column == "N000":
        return "Ruta relativa de la carpeta factory. Campo de control; no modificar."
    if column == "__subcarpeta":
        return "Subcarpeta original dentro de factory. Campo de control; no modificar."
    if column == "__archivo_origen":
        return "Nombre original del fichero. Campo de control; no modificar."
    if column == "__archivo_salida":
        return "Nombre del fichero .lparam generado. Puede editarse; las colisiones abortan la exportación."
    if column == "__id":
        return "Identificador interno. No modificar."
    base = column_base_key(column)
    if not base:
        return ""
    info = metadata.get(base, {})
    title = str(info.get("title") or "")
    status = str(info.get("status") or "")
    data_type = str(info.get("data_type") or "")
    parts = [p for p in [title, f"Metadata R/S: {status}" if status else "", f"Tipo: {data_type}" if data_type else ""] if p]
    if "#" in column:
        parts.append("Ocurrencia repetida del mismo parámetro en el fichero original.")
    return "\n".join(parts)


def create_workbook_for_factory(factory_dir: Path, process_root: Path, metadata: dict[str, dict[str, object]], output_path: Path | None = None) -> tuple[Path, int]:
    factory_dir = Path(factory_dir)
    process_root = Path(process_root)
    output_path = Path(output_path) if output_path else factory_dir / "parametros.xlsx"
    files = find_parameter_files(factory_dir)
    if not files:
        raise ParamEditorError(f"No hay ficheros de parámetros en {factory_dir}")

    documents: list[tuple[Path, ParamDocument]] = [(p, ParamDocument.from_path(p)) for p in files]
    max_occurrences: dict[str, int] = defaultdict(int)
    for _, doc in documents:
        counts = Counter(entry.key for entry in doc.entries)
        for key, count in counts.items():
            max_occurrences[key] = max(max_occurrences[key], count)

    all_keys = set(metadata.keys()) | set(max_occurrences.keys())
    parameter_columns: list[str] = []
    for key in sorted(all_keys, key=lambda x: int(x[1:]) if re.fullmatch(r"N\d{3}", x) else 999999):
        occurrences = max(1, max_occurrences.get(key, 0))
        for occurrence in range(1, occurrences + 1):
            parameter_columns.append(key if occurrence == 1 else f"{key}#{occurrence}")

    headers = SPECIAL_HEADERS + [c for c in parameter_columns if c != "N000"]

    wb = Workbook()
    ws = wb.active
    ws.title = VISIBLE_SHEET
    source_ws = wb.create_sheet(SOURCE_SHEET)
    source_ws.sheet_state = "veryHidden"

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    control_fill = PatternFill("solid", fgColor="E7E6E6")
    header_font = Font(bold=True)
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(1, col_idx, header)
        cell.font = header_font
        cell.fill = control_fill if header.startswith("__") or header == "N000" else header_fill
        comment = _header_comment(header, metadata)
        if comment:
            cell.comment = Comment(comment, "Param_editor")

    source_headers = ["id", "factory_rel", "source_rel", "original_name", "raw_b64", "encoding"]
    source_ws.append(source_headers)

    factory_rel = factory_dir.relative_to(process_root).as_posix()
    for row_idx, (file_path, doc) in enumerate(documents, 2):
        recipe_id = f"R{row_idx-1:06d}"
        source_rel = file_path.relative_to(factory_dir).as_posix()
        subfolder = file_path.parent.relative_to(factory_dir).as_posix()
        if subfolder == ".":
            subfolder = ""
        output_name = file_path.name if file_path.name.lower().endswith(".lparam") else f"{file_path.name}.lparam"
        value_map = doc.value_map()

        row_values = {
            "N000": factory_rel,
            "__subcarpeta": subfolder,
            "__archivo_origen": file_path.name,
            "__archivo_salida": output_name,
            "__id": recipe_id,
            **value_map,
        }
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row_idx, col_idx, row_values.get(header, ""))
            cell.number_format = "@"

        source_ws.append([
            recipe_id,
            factory_rel,
            source_rel,
            file_path.name,
            base64.b64encode(doc.raw_bytes).decode("ascii"),
            doc.encoding,
        ])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    # Control columns
    for idx, header in enumerate(headers, 1):
        if header == "__id":
            ws.column_dimensions[get_column_letter(idx)].hidden = True
        elif header in {"N000", "__subcarpeta", "__archivo_origen", "__archivo_salida"}:
            ws.column_dimensions[get_column_letter(idx)].width = 28
        else:
            ws.column_dimensions[get_column_letter(idx)].width = 14

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path, len(documents)


def read_source_records(wb) -> dict[str, dict[str, str]]:
    if SOURCE_SHEET not in wb.sheetnames:
        raise ParamEditorError(
            "El Excel no contiene la hoja interna _SOURCE. Debe regenerarse con PARAMtoEXCEL para garantizar un round-trip seguro."
        )
    ws = wb[SOURCE_SHEET]
    headers = [str(c.value) if c.value is not None else "" for c in ws[1]]
    records: dict[str, dict[str, str]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = {headers[i]: ("" if row[i] is None else str(row[i])) for i in range(min(len(headers), len(row)))}
        if record.get("id"):
            records[record["id"]] = record
    return records


@dataclass
class PlannedOutput:
    target: Path
    source_xlsx: Path
    source_rel: str
    content: bytes
    changed_columns: list[str]


def build_export_plan(xlsx_path: Path, process_root: Path, output_root: Path, metadata: dict[str, dict[str, object]]) -> list[PlannedOutput]:
    xlsx_path = Path(xlsx_path)
    wb = load_workbook(xlsx_path, data_only=False)
    if VISIBLE_SHEET not in wb.sheetnames:
        raise ParamEditorError(f"{xlsx_path}: no existe la hoja {VISIBLE_SHEET}.")
    ws = wb[VISIBLE_SHEET]
    records = read_source_records(wb)
    headers = [("" if c.value is None else str(c.value).strip()) for c in ws[1]]
    header_idx = {h: i + 1 for i, h in enumerate(headers) if h}
    required = {"N000", "__subcarpeta", "__archivo_origen", "__archivo_salida", "__id"}
    missing = required - set(header_idx)
    if missing:
        raise ParamEditorError(f"{xlsx_path}: faltan columnas de control {sorted(missing)}")

    default_status = metadata_default_status(metadata)
    plan: list[PlannedOutput] = []
    errors: list[str] = []

    for row_idx in range(2, ws.max_row + 1):
        recipe_id_cell = ws.cell(row_idx, header_idx["__id"]).value
        if recipe_id_cell is None or str(recipe_id_cell).strip() == "":
            if all(ws.cell(row_idx, c).value in (None, "") for c in range(1, ws.max_column + 1)):
                continue
            errors.append(f"Fila {row_idx}: falta __id; no se puede identificar la receta original.")
            continue
        recipe_id = str(recipe_id_cell).strip()
        source = records.get(recipe_id)
        if not source:
            errors.append(f"Fila {row_idx}: __id {recipe_id!r} no existe en _SOURCE.")
            continue

        try:
            factory_rel = str(ws.cell(row_idx, header_idx["N000"]).value or "")
            subfolder = str(ws.cell(row_idx, header_idx["__subcarpeta"]).value or "")
            original_name = str(ws.cell(row_idx, header_idx["__archivo_origen"]).value or "")
            output_name = normalize_output_name(str(ws.cell(row_idx, header_idx["__archivo_salida"]).value or ""))

            # Campos de trazabilidad no deben alterarse.
            if factory_rel.replace("\\", "/") != source["factory_rel"].replace("\\", "/"):
                raise ParamEditorError("N000 fue modificado; la ruta de origen es un campo de control.")
            expected_subfolder = str(Path(source["source_rel"]).parent.as_posix())
            if expected_subfolder == ".":
                expected_subfolder = ""
            if subfolder.replace("\\", "/") != expected_subfolder.replace("\\", "/"):
                raise ParamEditorError("__subcarpeta fue modificada; es un campo de control.")
            if original_name != source["original_name"]:
                raise ParamEditorError("__archivo_origen fue modificado; es un campo de control.")

            raw = base64.b64decode(source["raw_b64"])
            doc = ParamDocument.from_bytes(raw)
            originals = doc.value_map()
            edits: dict[str, str] = {}
            changed_columns: list[str] = []
            for header, col_idx in header_idx.items():
                if not COLUMN_RE.match(header) or header == "N000":
                    continue
                cell_value = ws.cell(row_idx, col_idx).value
                new_value = "" if cell_value is None else str(cell_value)
                old_value = originals.get(header, "")
                if new_value != old_value:
                    base_key = column_base_key(header)
                    for msg in validate_edit(base_key or header, new_value, metadata):
                        errors.append(f"{xlsx_path.name}, fila {row_idx}, {header}: {msg}")
                    edits[header] = new_value
                    changed_columns.append(header)

            content = doc.render(edits, default_status)
            factory_path = safe_relative_path(source["factory_rel"], "factory_rel")
            subfolder_path = Path()
            if expected_subfolder:
                subfolder_path = safe_relative_path(expected_subfolder, "subcarpeta")
            target = Path(output_root) / factory_path / subfolder_path / output_name
            plan.append(PlannedOutput(target, xlsx_path, source["source_rel"], content, changed_columns))
        except ParamEditorError as exc:
            errors.append(f"{xlsx_path.name}, fila {row_idx}: {exc}")

    if errors:
        raise ParamEditorError("\n".join(errors))
    return plan


def validate_collisions(plan: Iterable[PlannedOutput]) -> None:
    by_target: dict[str, list[PlannedOutput]] = defaultdict(list)
    for item in plan:
        # Windows es case-insensitive: se valida así aunque la prueba corra en Linux.
        key = str(item.target).replace("\\", "/").casefold()
        by_target[key].append(item)
    collisions = [items for items in by_target.values() if len(items) > 1]
    if collisions:
        lines = ["Se han detectado colisiones de archivos de salida. No se ha escrito ningún LPARAM:"]
        for items in collisions:
            lines.append(f"  Destino: {items[0].target}")
            for item in items:
                lines.append(f"    - {item.source_xlsx}: {item.source_rel}")
        raise ParamEditorError("\n".join(lines))


def write_export_plan(plan: list[PlannedOutput], output_root: Path, manifest_name: str = "manifest.csv") -> Path:
    validate_collisions(plan)
    output_root = Path(output_root)
    output_root.mkdir(parents=True, exist_ok=True)

    for item in plan:
        item.target.parent.mkdir(parents=True, exist_ok=True)
        fd, tmp_name = tempfile.mkstemp(prefix=item.target.name + ".", suffix=".tmp", dir=str(item.target.parent))
        try:
            with os.fdopen(fd, "wb") as f:
                f.write(item.content)
            os.replace(tmp_name, item.target)
        except Exception:
            try:
                os.unlink(tmp_name)
            except FileNotFoundError:
                pass
            raise

    manifest = output_root / manifest_name
    with manifest.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["source_excel", "source_param", "output_file", "changed_columns"])
        for item in plan:
            writer.writerow([
                str(item.source_xlsx),
                item.source_rel,
                str(item.target),
                ";".join(item.changed_columns),
            ])
    return manifest


def discover_workbooks(process_root: Path) -> list[Path]:
    workbooks: list[Path] = []
    for factory in find_factory_dirs(process_root):
        candidate = factory / "parametros.xlsx"
        if candidate.exists():
            workbooks.append(candidate)
    return sorted(workbooks, key=lambda p: str(p).casefold())


def copy_process_fixture(source: Path, target: Path) -> None:
    if target.exists():
        shutil.rmtree(target)
    shutil.copytree(source, target)


def timestamp() -> str:
    return datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
